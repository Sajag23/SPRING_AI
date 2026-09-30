import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_exact_winning_technical_approach():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Colors
    BG_COLOR = RGBColor(255, 255, 255)
    WHITE = RGBColor(255, 255, 255)
    NAVY_TITLE = RGBColor(24, 43, 73)        # #182B49
    BLUE_TEXT = RGBColor(2, 132, 199)        # #0284C7
    DARK_BLUE_BOX = RGBColor(16, 57, 105)    # #103969
    LIGHT_ORANGE_PILL = RGBColor(254, 215, 170) # #FED7AA
    LIGHT_BLUE_PILL = RGBColor(186, 230, 253)   # #BAE6FD
    LIGHT_GREEN_PILL = RGBColor(187, 247, 208)  # #BBF7D0
    LIGHT_PURPLE_PILL = RGBColor(233, 213, 255) # #E9D5FF
    TEXT_BLACK = RGBColor(15, 23, 42)        # #0F172A
    GREEN_CHECK = RGBColor(22, 163, 74)      # #16A34A
    RED_CROSS = RGBColor(220, 38, 38)        # #DC2626
    FOOTER_BLUE = RGBColor(2, 119, 189)       # #0277BD
    BORDER_GREY = RGBColor(203, 213, 225)    # #CBD5E1

    # 1. Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()

    # 2. Header
    # Left Brand Title
    logo_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(2.2), Inches(0.6))
    tf_l = logo_box.text_frame
    p_l = tf_l.paragraphs[0]
    p_l.text = "SPRING AI"
    p_l.font.name = 'Georgia'
    p_l.font.size = Pt(22)
    p_l.font.bold = True
    p_l.font.color.rgb = RGBColor(217, 119, 6)

    # Center Header Title
    title_box = slide.shapes.add_textbox(Inches(2.7), Inches(0.15), Inches(7.5), Inches(0.6))
    tf_t = title_box.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = "TECHNICAL APPROACH"
    p_t.font.name = 'Georgia'
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = NAVY_TITLE
    p_t.alignment = PP_ALIGN.CENTER

    # Right SIH Badge
    sih_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.1), Inches(2.4), Inches(0.7))
    tf_s = sih_box.text_frame
    p_s1 = tf_s.paragraphs[0]
    p_s1.text = "SMART INDIA HACKATHON"
    p_s1.font.name = 'Arial'
    p_s1.font.size = Pt(9)
    p_s1.font.bold = True
    p_s1.font.color.rgb = NAVY_TITLE
    p_s1.alignment = PP_ALIGN.RIGHT

    p_s2 = tf_s.add_paragraph()
    p_s2.text = "2026"
    p_s2.font.name = 'Arial'
    p_s2.font.size = Pt(13)
    p_s2.font.bold = True
    p_s2.font.color.rgb = BLUE_TEXT
    p_s2.alignment = PP_ALIGN.RIGHT

    # -------------------------------------------------------------------------
    # 3. LEFT COLUMN: FLOW OF OUR SOLUTION
    # -------------------------------------------------------------------------
    lx = Inches(0.4)
    lw = Inches(3.4)

    # Flow Pill Header
    flow_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lx, Inches(0.85), lw, Inches(0.45))
    flow_pill.fill.solid()
    flow_pill.fill.fore_color.rgb = LIGHT_ORANGE_PILL
    flow_pill.line.fill.background()
    tf_fp = flow_pill.text_frame
    p_fp = tf_fp.paragraphs[0]
    p_fp.text = "FLOW OF OUR SOLUTION"
    p_fp.font.name = 'Georgia'
    p_fp.font.size = Pt(12)
    p_fp.font.bold = True
    p_fp.font.color.rgb = NAVY_TITLE
    p_fp.alignment = PP_ALIGN.CENTER

    # 6 Flow Step Nodes
    flow_steps = [
        ("Multi-Modal Data Integration", "(Copernicus DEM, IMD Rainfall, GSI Lithology, Sentinel-2)"),
        ("Catchment Feature Extraction", "(DEM Flow Accumulation, Slope & Drainage Density)"),
        ("PRSI Index Calculation", "(Automated Multi-Criteria Hydro-Geological Scoring)"),
        ("Automated Hazard & Safety Mask", "(Slope >30° Exclusion & Eco-Sensitive Filtering)"),
        ("Structure Recommendation", "(Staggered Trenches, Percolation Pits, Check Dams)"),
        ("1-Click MGNREGA Costing & DPR", "(PDF Dossier, GIS Dashboard & Alerts)")
    ]

    fy_start = Inches(1.4)
    node_h = Inches(0.82)
    node_gap = Inches(0.12)

    for idx, (step_title, step_desc) in enumerate(flow_steps):
        ny = fy_start + idx * (node_h + node_gap)
        
        # Node Box
        node_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lx, ny, lw, node_h)
        node_box.fill.solid()
        node_box.fill.fore_color.rgb = DARK_BLUE_BOX
        node_box.line.color.rgb = BLUE_TEXT
        node_box.line.width = Pt(1.5)

        tf_node = node_box.text_frame
        tf_node.word_wrap = True
        
        p_nt = tf_node.paragraphs[0]
        p_nt.text = step_title
        p_nt.font.name = 'Arial'
        p_nt.font.size = Pt(10)
        p_nt.font.bold = True
        p_nt.font.color.rgb = WHITE
        p_nt.alignment = PP_ALIGN.CENTER

        p_nd = tf_node.add_paragraph()
        p_nd.text = step_desc
        p_nd.font.name = 'Arial'
        p_nd.font.size = Pt(8)
        p_nd.font.color.rgb = LIGHT_BLUE_PILL
        p_nd.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------------------
    # 4. CENTER TOP: TECH STACK
    # -------------------------------------------------------------------------
    cx = Inches(4.1)
    cw = Inches(4.5)

    tech_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(0.85), cw, Inches(0.45))
    tech_pill.fill.solid()
    tech_pill.fill.fore_color.rgb = LIGHT_BLUE_PILL
    tech_pill.line.fill.background()
    tf_tp = tech_pill.text_frame
    p_tp = tf_tp.paragraphs[0]
    p_tp.text = "TECH STACK"
    p_tp.font.name = 'Georgia'
    p_tp.font.size = Pt(12)
    p_tp.font.bold = True
    p_tp.font.color.rgb = NAVY_TITLE
    p_tp.alignment = PP_ALIGN.CENTER

    # Outer Tech Stack Container Box
    tech_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.4), cw, Inches(3.0))
    tech_container.fill.solid()
    tech_container.fill.fore_color.rgb = RGBColor(248, 250, 252)
    tech_container.line.color.rgb = BORDER_GREY
    tech_container.line.width = Pt(1.5)

    tf_tc = tech_container.text_frame
    tf_tc.word_wrap = True

    tech_sections = [
        ("Deployment & Cloud", "Docker, Vercel, GCP / AWS Serverless Functions"),
        ("Security & Auth", "OAuth 2.0, JWT, Role-Based Access Control (RBAC)"),
        ("Frontend & GIS UI", "React.js, Mapbox GL JS, Leaflet.js, Tailwind CSS, Recharts"),
        ("Backend & Spatial DB", "Python (FastAPI / Flask), PostGIS, PyTorch, GDAL / Rasterio")
    ]

    for s_idx, (sec_title, sec_techs) in enumerate(tech_sections):
        p_st = tf_tc.paragraphs[0] if s_idx == 0 else tf_tc.add_paragraph()
        p_st.text = f"• {sec_title}:"
        p_st.font.name = 'Arial'
        p_st.font.size = Pt(9.5)
        p_st.font.bold = True
        p_st.font.color.rgb = NAVY_TITLE
        p_st.space_before = Pt(4) if s_idx > 0 else Pt(0)

        p_te = tf_tc.add_paragraph()
        p_te.text = f"   {sec_techs}"
        p_te.font.name = 'Arial'
        p_te.font.size = Pt(9)
        p_te.font.color.rgb = BLUE_TEXT
        p_te.space_after = Pt(4)

    # -------------------------------------------------------------------------
    # 5. RIGHT TOP: PRSI FORMULA CALCULATION
    # -------------------------------------------------------------------------
    rx = Inches(8.8)
    rw = Inches(4.133)

    prsi_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, Inches(0.85), rw, Inches(0.45))
    prsi_pill.fill.solid()
    prsi_pill.fill.fore_color.rgb = LIGHT_GREEN_PILL
    prsi_pill.line.fill.background()
    tf_prp = prsi_pill.text_frame
    p_prp = tf_prp.paragraphs[0]
    p_prp.text = "PRSI FORMULA CALCULATION"
    p_prp.font.name = 'Georgia'
    p_prp.font.size = Pt(12)
    p_prp.font.bold = True
    p_prp.font.color.rgb = NAVY_TITLE
    p_prp.alignment = PP_ALIGN.CENTER

    prsi_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, Inches(1.4), rw, Inches(3.0))
    prsi_container.fill.solid()
    prsi_container.fill.fore_color.rgb = RGBColor(248, 250, 252)
    prsi_container.line.color.rgb = GREEN_CHECK
    prsi_container.line.width = Pt(1.5)

    tf_pc = prsi_container.text_frame
    tf_pc.word_wrap = True

    p_f1 = tf_pc.paragraphs[0]
    p_f1.text = "❏ PRSI Weighted Hydro-Geological Formula:"
    p_f1.font.name = 'Arial'
    p_f1.font.size = Pt(9.5)
    p_f1.font.bold = True
    p_f1.font.color.rgb = NAVY_TITLE

    p_f2 = tf_pc.add_paragraph()
    p_f2.text = "PRSI = 0.35(S) + 0.25(L) + 0.20(G) + 0.15(R) + 0.05(D)"
    p_f2.font.name = 'Georgia'
    p_f2.font.size = Pt(10.5)
    p_f2.font.bold = True
    p_f2.font.color.rgb = BLUE_TEXT
    p_f2.alignment = PP_ALIGN.CENTER
    p_f2.space_before = Pt(4)
    p_f2.space_after = Pt(6)

    p_f3 = tf_pc.add_paragraph()
    p_f3.text = "Where Parameter Weights Are:"
    p_f3.font.name = 'Arial'
    p_f3.font.size = Pt(9)
    p_f3.font.bold = True
    p_f3.font.color.rgb = TEXT_BLACK

    weights = [
        "S = Slope Factor (35% weightage)",
        "L = Rock Lithology & Infiltration (25% weightage)",
        "G = Groundwater Depth (20% weightage)",
        "R = Rainfall Distribution (15% weightage)",
        "D = Drainage Density (5% weightage)"
    ]

    for w_text in weights:
        p_w = tf_pc.add_paragraph()
        p_w.text = f"  • {w_text}"
        p_w.font.name = 'Arial'
        p_w.font.size = Pt(8.5)
        p_w.font.color.rgb = TEXT_BLACK

    # -------------------------------------------------------------------------
    # 6. BOTTOM RIGHT: ROLE BASED ACCESS CONTROL (RBAC) MATRIX TABLE
    # -------------------------------------------------------------------------
    rb_x = Inches(4.1)
    rb_w = Inches(8.833)
    rb_y = Inches(4.55)

    rbac_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rb_x, rb_y, rb_w, Inches(0.45))
    rbac_pill.fill.solid()
    rbac_pill.fill.fore_color.rgb = LIGHT_PURPLE_PILL
    rbac_pill.line.fill.background()
    tf_rp = rbac_pill.text_frame
    p_rp = tf_rp.paragraphs[0]
    p_rp.text = "ROLE BASED ACCESS CONTROL MATRIX"
    p_rp.font.name = 'Georgia'
    p_rp.font.size = Pt(12)
    p_rp.font.bold = True
    p_rp.font.color.rgb = NAVY_TITLE
    p_rp.alignment = PP_ALIGN.CENTER

    # Table: 5 Rows x 9 Columns
    table_shape = slide.shapes.add_table(5, 9, rb_x, rb_y + Inches(0.5), rb_w, Inches(1.9))
    t = table_shape.table

    t.columns[0].width = Inches(1.833)
    for c in range(1, 9):
        t.columns[c].width = Inches(0.875)

    rbac_headers = ["ROLE", "GIS MAPS", "PRSI MODEL", "REPORTS/DPR", "GW DATA", "FORECAST", "ALERTS", "OFFLINE SYNC", "RESEARCH"]
    rbac_rows = [
        ["MINISTRY OFFICER", "✔", "✔", "✔", "✔", "✔", "✔", "❌", "✔"],
        ["HYDROLOGIST / SCIENTIST", "✔", "✔", "✔", "✔", "✔", "✔", "❌", "✔"],
        ["DISTRICT OFFICER", "✔", "✔", "✔", "✔", "✔", "✔", "✔", "❌"],
        ["FIELD SURVEYOR (JAL SAHI)", "✔", "❌", "✔", "❌", "❌", "✔", "✔", "❌"]
    ]

    # Set Headers
    for c_i, h_text in enumerate(rbac_headers):
        cell = t.cell(0, c_i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE_BOX
        p_c = cell.text_frame.paragraphs[0]
        p_c.text = h_text
        p_c.font.name = 'Arial'
        p_c.font.size = Pt(8)
        p_c.font.bold = True
        p_c.font.color.rgb = WHITE
        p_c.alignment = PP_ALIGN.CENTER

    # Set Row Values
    for r_i, r_vals in enumerate(rbac_rows, start=1):
        for c_i, val in enumerate(r_vals):
            cell = t.cell(r_i, c_i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(248, 250, 252) if r_i % 2 == 1 else WHITE
            
            p_c = cell.text_frame.paragraphs[0]
            p_c.text = val
            p_c.font.name = 'Arial'
            p_c.font.size = Pt(8.5)
            p_c.font.bold = True
            if val == "✔":
                p_c.font.color.rgb = GREEN_CHECK
            elif val == "❌":
                p_c.font.color.rgb = RED_CROSS
            else:
                p_c.font.color.rgb = NAVY_TITLE
                p_c.alignment = PP_ALIGN.LEFT
            if c_i > 0:
                p_c.alignment = PP_ALIGN.CENTER

    # 7. Footer Blue Bar
    footer_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.45))
    footer_bar.fill.solid()
    footer_bar.fill.fore_color.rgb = FOOTER_BLUE
    footer_bar.line.fill.background()

    tf_foot = footer_bar.text_frame
    p_foot_l = tf_foot.paragraphs[0]
    p_foot_l.text = "@SIH Idea submission- Template"
    p_foot_l.font.name = 'Arial'
    p_foot_l.font.size = Pt(11)
    p_foot_l.font.color.rgb = WHITE
    p_foot_l.alignment = PP_ALIGN.LEFT

    num_box = slide.shapes.add_textbox(Inches(12.2), Inches(7.05), Inches(1.0), Inches(0.45))
    tf_num = num_box.text_frame
    p_num = tf_num.paragraphs[0]
    p_num.text = "3"
    p_num.font.name = 'Arial'
    p_num.font.size = Pt(12)
    p_num.font.bold = True
    p_num.font.color.rgb = WHITE
    p_num.alignment = PP_ALIGN.RIGHT

    out_path = "SPRING_AI_Exact_Winning_Technical_Approach.pptx"
    prs.save(out_path)
    print(f"Successfully generated exact winning technical approach slide at: {os.path.abspath(out_path)}")

if __name__ == "__main__":
    build_exact_winning_technical_approach()
