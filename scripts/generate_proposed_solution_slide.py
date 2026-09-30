import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_proposed_solution_slide():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Use a blank slide layout
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # Color Palette
    BG_COLOR = RGBColor(248, 250, 252)       # #F8FAFC
    NAVY = RGBColor(15, 23, 42)             # #0F172A
    PRIMARY_BLUE = RGBColor(37, 99, 235)     # #2563EB
    DARK_BLUE = RGBColor(2, 132, 199)       # #0284C7
    EMERALD = RGBColor(5, 150, 105)         # #059669
    AMBER = RGBColor(217, 119, 6)           # #D97706
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(71, 85, 105)       # #475569
    LIGHT_BORDER = RGBColor(226, 232, 240)  # #E2E8F0
    FOOTER_BLUE = RGBColor(2, 119, 189)     # #0277BD

    # 1. Background Fill
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_COLOR
    bg_shape.line.fill.background()

    # 2. Header Elements
    # Team Name Oval (Top Left)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.4), Inches(0.3), Inches(1.8), Inches(0.8))
    oval.fill.solid()
    oval.fill.fore_color.rgb = WHITE
    oval.line.color.rgb = PRIMARY_BLUE
    oval.line.width = Pt(1.5)
    tf_oval = oval.text_frame
    tf_oval.word_wrap = True
    p_oval = tf_oval.paragraphs[0]
    p_oval.text = "Your Team Name"
    p_oval.font.name = 'Arial'
    p_oval.font.size = Pt(11)
    p_oval.font.color.rgb = NAVY
    p_oval.alignment = PP_ALIGN.CENTER

    # Center Header Title: IDEA TITLE (SPRING AI)
    title_box = slide.shapes.add_textbox(Inches(3.5), Inches(0.25), Inches(6.333), Inches(0.8))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = "SPRING AI"
    p_title.font.name = 'Georgia'
    p_title.font.size = Pt(30)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY
    p_title.alignment = PP_ALIGN.CENTER

    p_sub = tf_title.add_paragraph()
    p_sub.text = "AI-Driven Springshed Rejuvenation & Hydro-Geological Planning System"
    p_sub.font.name = 'Arial'
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = GRAY_TEXT
    p_sub.alignment = PP_ALIGN.CENTER

    # Top Right SIH Logo Placeholder / Badge
    sih_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.6), Inches(0.25), Inches(2.3), Inches(0.85))
    sih_box.fill.solid()
    sih_box.fill.fore_color.rgb = WHITE
    sih_box.line.color.rgb = LIGHT_BORDER
    tf_sih = sih_box.text_frame
    p_sih1 = tf_sih.paragraphs[0]
    p_sih1.text = "SMART INDIA HACKATHON"
    p_sih1.font.name = 'Arial'
    p_sih1.font.size = Pt(9.5)
    p_sih1.font.bold = True
    p_sih1.font.color.rgb = NAVY
    p_sih1.alignment = PP_ALIGN.CENTER

    p_sih2 = tf_sih.add_paragraph()
    p_sih2.text = "2026"
    p_sih2.font.name = 'Arial'
    p_sih2.font.size = Pt(14)
    p_sih2.font.bold = True
    p_sih2.font.color.rgb = PRIMARY_BLUE
    p_sih2.alignment = PP_ALIGN.CENTER

    # 3. Main Slide Section Title
    main_title_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.2), Inches(12.5), Inches(0.6))
    tf_mt = main_title_box.text_frame
    p_mt = tf_mt.paragraphs[0]
    p_mt.text = "• Proposed Solution (Describe your Idea/Solution/Prototype)"
    p_mt.font.name = 'Arial'
    p_mt.font.size = Pt(20)
    p_mt.font.bold = True
    p_mt.font.color.rgb = PRIMARY_BLUE
    p_mt.font.underline = True

    # 4. Three High-Impact Visual Cards Layout
    card_width = Inches(3.9)
    card_height = Inches(4.9)
    card_y = Inches(1.85)
    card_xs = [Inches(0.4), Inches(4.7), Inches(9.0)]

    cards_data = [
        {
            "header_title": "Detailed Explanation",
            "header_icon": "💡",
            "accent_color": PRIMARY_BLUE,
            "bullets": [
                ("Multi-Modal Data Fusion", "Fuses Copernicus DEM, IMD Rainfall, GSI Lithology & Sentinel-2 LULC."),
                ("30-Sec DEM Delineation", "Automates 3D catchment boundary & recharge area mapping."),
                ("PRSI Hydro-Geological Engine", "Scores pixel suitability for groundwater recharge & structures."),
                ("Offline App & 1-Click DPR", "Field verification GIS app generating MGNREGA dossiers.")
            ],
            "pill": "⚡ 30-Sec Delineation Engine"
        },
        {
            "header_title": "Problem Addressed",
            "header_icon": "🎯",
            "accent_color": EMERALD,
            "bullets": [
                ("Eliminates Survey Delays", "Cuts 6-month manual ground mapping down to instantaneous AI."),
                ("Prevents Structure Failure", "Uses subsurface lithology & soil infiltration instead of surface-only."),
                ("Landslide Safety Exclusions", "Automated safety mask excludes steep terrain (>30° slope)."),
                ("Secures Tribal Water", "Extends spring baseflow (+45 days) across 100+ Tribal Districts.")
            ],
            "pill": "🌊 +45 Days Baseflow Extension"
        },
        {
            "header_title": "Innovation & Uniqueness",
            "header_icon": "⚡",
            "accent_color": AMBER,
            "bullets": [
                ("Proprietary PRSI Formula", "Multi-criteria hydro-geological decision matrix (Slope+Lithology+GW)."),
                ("Automated Safety Masking", "Hardcoded hazard & eco-sensitive constraint filtering."),
                ("Zero-Trust Offline Sync", "Enables remote field workers to capture data without internet."),
                ("Satellite-to-DPR Pipeline", "End-to-end automation from raw satellite DEM to MGNREGA budget.")
            ],
            "pill": "🚀 First-in-India GeoAI Matrix"
        }
    ]

    for i, data in enumerate(cards_data):
        cx = card_xs[i]

        # Card Outer Box Background
        card_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_width, card_height)
        card_bg.fill.solid()
        card_bg.fill.fore_color.rgb = WHITE
        card_bg.line.color.rgb = LIGHT_BORDER
        card_bg.line.width = Pt(1.5)

        # Header Banner
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, card_y, card_width, Inches(0.75))
        banner.fill.solid()
        banner.fill.fore_color.rgb = data["accent_color"]
        banner.line.fill.background()

        tf_b = banner.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = f"{data['header_icon']}  {data['header_title']}"
        p_b.font.name = 'Arial'
        p_b.font.size = Pt(14)
        p_b.font.bold = True
        p_b.font.color.rgb = WHITE
        p_b.alignment = PP_ALIGN.CENTER

        # Text Frame for Bullets
        content_box = slide.shapes.add_textbox(cx + Inches(0.15), card_y + Inches(0.85), card_width - Inches(0.3), Inches(3.3))
        tf_c = content_box.text_frame
        tf_c.word_wrap = True

        for b_idx, (b_head, b_desc) in enumerate(data["bullets"]):
            p_head = tf_c.paragraphs[0] if b_idx == 0 else tf_c.add_paragraph()
            p_head.text = f"• {b_head}"
            p_head.font.name = 'Arial'
            p_head.font.size = Pt(11)
            p_head.font.bold = True
            p_head.font.color.rgb = NAVY
            p_head.space_after = Pt(2)
            if b_idx > 0:
                p_head.space_before = Pt(8)

            p_desc = tf_c.add_paragraph()
            p_desc.text = f"   {b_desc}"
            p_desc.font.name = 'Arial'
            p_desc.font.size = Pt(9.5)
            p_desc.font.color.rgb = GRAY_TEXT
            p_desc.space_after = Pt(2)

        # Bottom Pill Badge
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.3), card_y + Inches(4.3), card_width - Inches(0.6), Inches(0.45))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(241, 245, 249) # #F1F5F9
        pill.line.color.rgb = data["accent_color"]
        pill.line.width = Pt(1)

        tf_p = pill.text_frame
        p_p = tf_p.paragraphs[0]
        p_p.text = data["pill"]
        p_p.font.name = 'Arial'
        p_p.font.size = Pt(10)
        p_p.font.bold = True
        p_p.font.color.rgb = data["accent_color"]
        p_p.alignment = PP_ALIGN.CENTER

    # 5. Footer Blue Bar
    footer_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.45))
    footer_bar.fill.solid()
    footer_bar.fill.fore_color.rgb = FOOTER_BLUE
    footer_bar.line.fill.background()

    tf_foot = footer_bar.text_frame
    p_foot_l = tf_foot.paragraphs[0]
    p_foot_l.text = "@SIH Idea submission- Template"
    p_foot_l.font.name = 'Arial'
    p_foot_l.font.size = Pt(11)
    p_foot_l.font.color.rgb = WHITE
    p_foot_l.alignment = PP_ALIGN.LEFT

    # Add page number on the right side of footer
    num_box = slide.shapes.add_textbox(Inches(12.2), Inches(7.05), Inches(1.0), Inches(0.45))
    tf_num = num_box.text_frame
    p_num = tf_num.paragraphs[0]
    p_num.text = "2"
    p_num.font.name = 'Arial'
    p_num.font.size = Pt(12)
    p_num.font.bold = True
    p_num.font.color.rgb = WHITE
    p_num.alignment = PP_ALIGN.RIGHT

    out_path = "SPRING_AI_Proposed_Solution_Slide.pptx"
    prs.save(out_path)
    print(f"Successfully generated Proposed Solution slide at: {os.path.abspath(out_path)}")

if __name__ == "__main__":
    build_proposed_solution_slide()
