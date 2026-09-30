import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_full_master_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(255, 255, 255)
    WHITE = RGBColor(255, 255, 255)
    NAVY_TITLE = RGBColor(24, 43, 73)
    BLUE_TEXT = RGBColor(2, 132, 199)
    DARK_BLUE_BOX = RGBColor(16, 57, 105)
    FOOTER_BLUE = RGBColor(2, 119, 189)
    BORDER_GREY = RGBColor(203, 213, 225)
    GREEN_CHECK = RGBColor(22, 163, 74)
    RED_CROSS = RGBColor(220, 38, 38)
    LIGHT_BLUE_BANNER = RGBColor(224, 242, 254)
    LIGHT_GREEN_PILL = RGBColor(220, 252, 231)
    PILL_BORDER = RGBColor(134, 239, 172)
    GRAY_TEXT = RGBColor(71, 85, 105)
    TEXT_BLACK = RGBColor(15, 23, 42)
    EMERALD = RGBColor(5, 150, 105)
    AMBER = RGBColor(217, 119, 6)
    PURPLE = RGBColor(124, 58, 237)

    def add_header(slide, title_text):
        oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.4), Inches(0.2), Inches(1.8), Inches(0.65))
        oval.fill.solid()
        oval.fill.fore_color.rgb = WHITE
        oval.line.color.rgb = BLUE_TEXT
        oval.line.width = Pt(1.5)
        tf_ov = oval.text_frame
        p_ov = tf_ov.paragraphs[0]
        p_ov.text = "SPRING AI"
        p_ov.font.name = 'Georgia'
        p_ov.font.size = Pt(11)
        p_ov.font.bold = True
        p_ov.font.color.rgb = NAVY_TITLE
        p_ov.alignment = PP_ALIGN.CENTER

        tbox = slide.shapes.add_textbox(Inches(2.5), Inches(0.15), Inches(7.8), Inches(0.6))
        tf_t = tbox.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = 'Georgia'
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_TITLE
        p_t.alignment = PP_ALIGN.CENTER

        sih_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.1), Inches(2.4), Inches(0.7))
        tf_s = sih_box.text_frame
        p_s1 = tf_s.paragraphs[0]
        p_s1.text = "SMART INDIA HACKATHON"
        p_s1.font.name = 'Arial'
        p_s1.font.size = Pt(9)
        p_s1.font.bold = True
        p_s1.font.color.rgb = NAVY_TITLE
        p_s1.alignment = PP_ALIGN.RIGHT

        p_s2 = tf_s.add_paragraph()
        p_s2.text = "2026"
        p_s2.font.name = 'Arial'
        p_s2.font.size = Pt(13)
        p_s2.font.bold = True
        p_s2.font.color.rgb = BLUE_TEXT
        p_s2.alignment = PP_ALIGN.RIGHT

    def add_footer(slide, page_num):
        footer_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.45))
        footer_bar.fill.solid()
        footer_bar.fill.fore_color.rgb = FOOTER_BLUE
        footer_bar.line.fill.background()

        tf_foot = footer_bar.text_frame
        p_foot_l = tf_foot.paragraphs[0]
        p_foot_l.text = "@SIH Idea submission- Template"
        p_foot_l.font.name = 'Arial'
        p_foot_l.font.size = Pt(11)
        p_foot_l.font.color.rgb = WHITE
        p_foot_l.alignment = PP_ALIGN.LEFT

        num_box = slide.shapes.add_textbox(Inches(12.2), Inches(7.05), Inches(1.0), Inches(0.45))
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = str(page_num)
        p_num.font.name = 'Arial'
        p_num.font.size = Pt(12)
        p_num.font.bold = True
        p_num.font.color.rgb = WHITE
        p_num.alignment = PP_ALIGN.RIGHT

    # SLIDE 1: COVER
    s1 = prs.slides.add_slide(blank_layout)
    add_header(s1, "SPRING AI — IDEA TITLE")
    hero = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(11.333), Inches(5.0))
    hero.fill.solid()
    hero.fill.fore_color.rgb = LIGHT_BLUE_BANNER
    hero.line.color.rgb = BLUE_TEXT
    hero.line.width = Pt(2)
    tf_h = hero.text_frame
    tf_h.word_wrap = True
    p_h1 = tf_h.paragraphs[0]
    p_h1.text = "SPRING AI"
    p_h1.font.name = 'Georgia'
    p_h1.font.size = Pt(36)
    p_h1.font.bold = True
    p_h1.font.color.rgb = NAVY_TITLE
    p_h1.alignment = PP_ALIGN.CENTER
    p_h2 = tf_h.add_paragraph()
    p_h2.text = "AI-Driven Springshed Rejuvenation & Hydro-Geological Planning System for Scheduled Tribal Regions"
    p_h2.font.name = 'Arial'
    p_h2.font.size = Pt(18)
    p_h2.font.bold = True
    p_h2.font.color.rgb = BLUE_TEXT
    p_h2.alignment = PP_ALIGN.CENTER
    p_h2.space_before = Pt(10)
    p_h2.space_after = Pt(20)
    p_h3 = tf_h.add_paragraph()
    p_h3.text = "• Problem Statement ID: 1729 | Ministry of Tribal Affairs (MoTA)\n• Target Region: 100+ Scheduled Tribal Districts across Western Ghats, Northeast & Central India\n• Key Innovation: Copernicus 30m DEM Flow Accumulation + PRSI Hydro-Geological Scoring + MGNREGA DPR Dossiers"
    p_h3.font.name = 'Arial'
    p_h3.font.size = Pt(14)
    p_h3.font.color.rgb = NAVY_TITLE
    p_h3.alignment = PP_ALIGN.CENTER
    add_footer(s1, 1)

    # Note: Save master deck
    out_master = "SPRING_AI_SIH_2026_MASTER_WINNING_DECK.pptx"
    prs.save(out_master)
    print(f"Master presentation compiled successfully at: {os.path.abspath(out_master)}")

if __name__ == "__main__":
    build_full_master_deck()
