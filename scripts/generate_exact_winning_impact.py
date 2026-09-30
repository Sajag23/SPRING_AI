import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_exact_winning_impact():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Colors
    BG_COLOR = RGBColor(255, 255, 255)
    WHITE = RGBColor(255, 255, 255)
    NAVY_TITLE = RGBColor(24, 43, 73)        # #182B49
    BLUE_TEXT = RGBColor(2, 132, 199)        # #0284C7
    DARK_BLUE_BOX = RGBColor(16, 57, 105)    # #103969
    LIGHT_BLUE_BANNER = RGBColor(224, 242, 254) # #E0F2FE
    LIGHT_BLUE_CARD = RGBColor(238, 242, 255)   # #EEF2FF
    BORDER_BLUE = RGBColor(186, 230, 253)       # #BAE6FD
    TEXT_BLACK = RGBColor(15, 23, 42)        # #0F172A
    GRAY_TEXT = RGBColor(71, 85, 105)        # #475569
    FOOTER_BLUE = RGBColor(2, 119, 189)       # #0277BD
    FOREST_GREEN = RGBColor(5, 150, 105)     # #059669

    # 1. Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()

    # 2. Header
    # Left Oval Logo Box
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.4), Inches(0.2), Inches(1.8), Inches(0.65))
    oval.fill.solid()
    oval.fill.fore_color.rgb = WHITE
    oval.line.color.rgb = BLUE_TEXT
    oval.line.width = Pt(1.5)
    tf_ov = oval.text_frame
    p_ov = tf_ov.paragraphs[0]
    p_ov.text = "SPRING AI"
    p_ov.font.name = 'Georgia'
    p_ov.font.size = Pt(11)
    p_ov.font.bold = True
    p_ov.font.color.rgb = NAVY_TITLE
    p_ov.alignment = PP_ALIGN.CENTER

    # Center Header Title
    title_box = slide.shapes.add_textbox(Inches(2.5), Inches(0.15), Inches(7.8), Inches(0.6))
    tf_t = title_box.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = "IMPACT AND BENEFITS"
    p_t.font.name = 'Georgia'
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = NAVY_TITLE
    p_t.alignment = PP_ALIGN.CENTER

    # Right SIH Badge
    sih_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.1), Inches(2.4), Inches(0.7))
    tf_s = sih_box.text_frame
    p_s1 = tf_s.paragraphs[0]
    p_s1.text = "SMART INDIA HACKATHON"
    p_s1.font.name = 'Arial'
    p_s1.font.size = Pt(9)
    p_s1.font.bold = True
    p_s1.font.color.rgb = NAVY_TITLE
    p_s1.alignment = PP_ALIGN.RIGHT

    p_s2 = tf_s.add_paragraph()
    p_s2.text = "2026"
    p_s2.font.name = 'Arial'
    p_s2.font.size = Pt(13)
    p_s2.font.bold = True
    p_s2.font.color.rgb = BLUE_TEXT
    p_s2.alignment = PP_ALIGN.RIGHT

    # 3. Top Executive Summary Banner
    summary_banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(0.85), Inches(9.333), Inches(0.65))
    summary_banner.fill.solid()
    summary_banner.fill.fore_color.rgb = LIGHT_BLUE_BANNER
    summary_banner.line.color.rgb = BLUE_TEXT
    summary_banner.line.width = Pt(1)

    tf_sb = summary_banner.text_frame
    p_sb = tf_sb.paragraphs[0]
    p_sb.text = "Solar panels boomed because of rising electricity costs + government push.\nSimilarly, acute water scarcity + spring rejuvenation mandates are pushing SPRING AI forward."
    p_sb.font.name = 'Arial'
    p_sb.font.size = Pt(10)
    p_sb.font.bold = True
    p_sb.font.color.rgb = NAVY_TITLE
    p_sb.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------------------
    # 4. LEFT SECTION: ICEBERG GRAPHIC MODEL (OR EMBEDDED IMAGE)
    # -------------------------------------------------------------------------
    iceberg_path = "assets/iceberg_graphic.png"
    if os.path.exists(iceberg_path):
        slide.shapes.add_picture(iceberg_path, Inches(0.4), Inches(1.65), Inches(4.5), Inches(5.2))
    else:
        # Fallback Iceberg Container Box
        ib_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(1.65), Inches(4.5), Inches(5.2))
        ib_box.fill.solid()
        ib_box.fill.fore_color.rgb = RGBColor(240, 249, 255)
        ib_box.line.color.rgb = BLUE_TEXT
        tf_ib = ib_box.text_frame
        p_ib = tf_ib.paragraphs[0]
        p_ib.text = "ICEBERG ADOPTION MODEL\n\n• Surface: Drying Springs & Field Surveys\n\n• Submerged: 18.5 Billion Liters Untapped Recharge Potential & Slope >30° Safety Filters"
        p_ib.font.color.rgb = NAVY_TITLE

    # -------------------------------------------------------------------------
    # 5. RIGHT TOP: 5-TIER MULTI-LEVEL BENEFIT CARDS
    # -------------------------------------------------------------------------
    rx = Inches(5.1)
    rw = Inches(7.833)

    tiers_data = [
        ("Household Level", "Lower water bills, 4.2M Liters annual water security, easy self-assessment GIS tool.", Inches(5.1), Inches(1.65), Inches(3.7), Inches(1.0)),
        ("Community Level", "Collective water conservation, secures 100+ Scheduled Tribal Districts & livelihoods.", Inches(9.0), Inches(1.65), Inches(3.933), Inches(1.0)),
        ("Regional Level", "Groundwater table rise (+1.2m), baseflow extension (+45 days), reduced flash runoff.", Inches(5.1), Inches(2.75), Inches(3.7), Inches(1.0)),
        ("Government Level", "Data-driven policymaking, easier monitoring of recharge mandates, MGNREGA budget optimization.", Inches(9.0), Inches(2.75), Inches(3.933), Inches(1.0)),
        ("Environmental / Long-Term", "Groundwater replenishment, climate change resilience, sustainable future mountain water resources.", Inches(5.1), Inches(3.85), Inches(7.833), Inches(0.95))
    ]

    for title, desc, cx, cy, cw, ch in tiers_data:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch)
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BLUE_CARD
        card.line.color.rgb = BORDER_BLUE
        card.line.width = Pt(1.2)

        tf_c = card.text_frame
        tf_c.word_wrap = True

        p_t = tf_c.paragraphs[0]
        p_t.text = f"• {title}"
        p_t.font.name = 'Georgia'
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_TITLE

        p_d = tf_c.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.name = 'Arial'
        p_d.font.size = Pt(9)
        p_d.font.color.rgb = GRAY_TEXT

    # -------------------------------------------------------------------------
    # 6. RIGHT BOTTOM: SPRINGS HED REJUVENATION VALUE CHAIN
    # -------------------------------------------------------------------------
    vc_y = Inches(4.95)
    
    vc_title = slide.shapes.add_textbox(rx, vc_y, rw, Inches(0.35))
    tf_vt = vc_title.text_frame
    p_vt = tf_vt.paragraphs[0]
    p_vt.text = "Springshed Rejuvenation Value Chain"
    p_vt.font.name = 'Georgia'
    p_vt.font.size = Pt(13)
    p_vt.font.bold = True
    p_vt.font.color.rgb = NAVY_TITLE
    p_vt.alignment = PP_ALIGN.CENTER

    vc_steps = [
        ("Untapped Resource", "Springs Available in Mountain Belts"),
        ("AI Delineation", "Copernicus DEM & PRSI Scoring"),
        ("Field Sync", "Jal Sahi Mobile GIS App"),
        ("Govt / MGNREGA", "Itemized DPR Budget Sanction"),
        ("Sustainable Water", "+45 Days Extended Baseflow")
    ]

    vw = Inches(1.45)
    gap_v = Inches(0.12)
    start_vx = rx

    for idx, (v_head, v_desc) in enumerate(vc_steps):
        vx = start_vx + idx * (vw + gap_v)

        v_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, vx, vc_y + Inches(0.4), vw, Inches(1.5))
        v_box.fill.solid()
        v_box.fill.fore_color.rgb = WHITE
        v_box.line.color.rgb = BLUE_TEXT
        v_box.line.width = Pt(1.5)

        tf_v = v_box.text_frame
        tf_v.word_wrap = True

        # Diamond Icon Placeholder
        p_ic = tf_v.paragraphs[0]
        p_ic.text = "◆"
        p_ic.font.size = Pt(16)
        p_ic.font.color.rgb = FOREST_GREEN if idx % 2 == 0 else BLUE_TEXT
        p_ic.alignment = PP_ALIGN.CENTER

        p_vh = tf_v.add_paragraph()
        p_vh.text = v_head
        p_vh.font.name = 'Arial'
        p_vh.font.size = Pt(9)
        p_vh.font.bold = True
        p_vh.font.color.rgb = NAVY_TITLE
        p_vh.alignment = PP_ALIGN.CENTER
        p_vh.space_before = Pt(2)

        p_vd = tf_v.add_paragraph()
        p_vd.text = v_desc
        p_vd.font.name = 'Arial'
        p_vd.font.size = Pt(7.5)
        p_vd.font.color.rgb = GRAY_TEXT
        p_vd.alignment = PP_ALIGN.CENTER
        p_vd.space_before = Pt(2)

    # 7. Footer Blue Bar
    footer_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.45))
    footer_bar.fill.solid()
    footer_bar.fill.fore_color.rgb = FOOTER_BLUE
    footer_bar.line.fill.background()

    tf_foot = footer_bar.text_frame
    p_foot_l = tf_foot.paragraphs[0]
    p_foot_l.text = "@SIH Idea submission- Template"
    p_foot_l.font.name = 'Arial'
    p_foot_l.font.size = Pt(11)
    p_foot_l.font.color.rgb = WHITE
    p_foot_l.alignment = PP_ALIGN.LEFT

    num_box = slide.shapes.add_textbox(Inches(12.2), Inches(7.05), Inches(1.0), Inches(0.45))
    tf_num = num_box.text_frame
    p_num = tf_num.paragraphs[0]
    p_num.text = "5"
    p_num.font.name = 'Arial'
    p_num.font.size = Pt(12)
    p_num.font.bold = True
    p_num.font.color.rgb = WHITE
    p_num.alignment = PP_ALIGN.RIGHT

    out_path = "SPRING_AI_Exact_Winning_Impact.pptx"
    prs.save(out_path)
    print(f"Successfully generated exact winning Impact slide at: {os.path.abspath(out_path)}")

if __name__ == "__main__":
    build_exact_winning_impact()
