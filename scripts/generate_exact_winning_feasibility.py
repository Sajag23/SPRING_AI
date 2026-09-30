import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_exact_winning_feasibility():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Colors
    BG_COLOR = RGBColor(255, 255, 255)
    WHITE = RGBColor(255, 255, 255)
    NAVY_TITLE = RGBColor(24, 43, 73)        # #182B49
    BLUE_TEXT = RGBColor(2, 132, 199)        # #0284C7
    DARK_BLUE_TAG = RGBColor(16, 57, 105)    # #103969
    LIGHT_GREEN_PILL = RGBColor(220, 252, 231) # #DCFCE7
    PILL_BORDER = RGBColor(134, 239, 172)     # #86EFAC
    TEXT_BLACK = RGBColor(15, 23, 42)        # #0F172A
    GRAY_TEXT = RGBColor(71, 85, 105)        # #475569
    FOOTER_BLUE = RGBColor(2, 119, 189)       # #0277BD
    BORDER_GREY = RGBColor(203, 213, 225)    # #CBD5E1

    # 1. Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()

    # 2. Header
    # Left Oval Logo Box
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.4), Inches(0.2), Inches(1.8), Inches(0.65))
    oval.fill.solid()
    oval.fill.fore_color.rgb = WHITE
    oval.line.color.rgb = BLUE_TEXT
    oval.line.width = Pt(1.5)
    tf_ov = oval.text_frame
    p_ov = tf_ov.paragraphs[0]
    p_ov.text = "SPRING AI"
    p_ov.font.name = 'Georgia'
    p_ov.font.size = Pt(11)
    p_ov.font.bold = True
    p_ov.font.color.rgb = NAVY_TITLE
    p_ov.alignment = PP_ALIGN.CENTER

    # Center Header Title
    title_box = slide.shapes.add_textbox(Inches(2.5), Inches(0.15), Inches(7.8), Inches(0.6))
    tf_t = title_box.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = "FEASIBILITY AND VIABILITY"
    p_t.font.name = 'Georgia'
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = NAVY_TITLE
    p_t.alignment = PP_ALIGN.CENTER

    # Right SIH Badge
    sih_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.1), Inches(2.4), Inches(0.7))
    tf_s = sih_box.text_frame
    p_s1 = tf_s.paragraphs[0]
    p_s1.text = "SMART INDIA HACKATHON"
    p_s1.font.name = 'Arial'
    p_s1.font.size = Pt(9)
    p_s1.font.bold = True
    p_s1.font.color.rgb = NAVY_TITLE
    p_s1.alignment = PP_ALIGN.RIGHT

    p_s2 = tf_s.add_paragraph()
    p_s2.text = "2026"
    p_s2.font.name = 'Arial'
    p_s2.font.size = Pt(13)
    p_s2.font.bold = True
    p_s2.font.color.rgb = BLUE_TEXT
    p_s2.alignment = PP_ALIGN.RIGHT

    # -------------------------------------------------------------------------
    # 3. LEFT SECTION: FEASIBILITY, CHALLENGES, STRATEGY (CONTAINER BOX WITH NODES)
    # -------------------------------------------------------------------------
    lx = Inches(0.4)
    lw = Inches(8.2)
    lh = Inches(5.6)
    ly = Inches(0.95)

    # Outer Border Container
    outer_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lx, ly, lw, lh)
    outer_container.fill.solid()
    outer_container.fill.fore_color.rgb = RGBColor(248, 250, 252)
    outer_container.line.color.rgb = BLUE_TEXT
    outer_container.line.width = Pt(1.5)

    # Vertical Connecting Line behind numbered circles
    line_x = lx + Inches(0.45)
    conn_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, line_x - Inches(0.03), ly + Inches(0.5), Inches(0.06), Inches(4.5))
    conn_line.fill.solid()
    conn_line.fill.fore_color.rgb = BLUE_TEXT
    conn_line.line.fill.background()

    # 3 Sections Data
    sections_data = [
        {
            "num": "1", "title": "Feasibility",
            "bullets": [
                ("Technically → ", "Existing satellite APIs (Copernicus DEM, IMD, GSI), GIS mapping, and custom PRSI algorithms."),
                ("Economically → ", "Optimizes MGNREGA budgets, prevents structural failure, and saves millions in manual surveys."),
                ("Socially → ", "Solves drinking water scarcity across 100+ Scheduled Tribal Districts & aligns with Ministry goals.")
            ]
        },
        {
            "num": "2", "title": "Challenges",
            "bullets": [
                ("Data Accuracy → ", "Remote tribal mountain terrain data & local rainfall observations may be sparse."),
                ("User Adoption → ", "Low digital literacy among Gram Panchayat field surveyors (Jal Sahi)."),
                ("Ground Truth Reliability → ", "Risk of unverified field reports or incorrect GPS site coordinates."),
                ("Regulatory Barriers → ", "Compliance with steep terrain safety (>30° slope) & forest conservation laws.")
            ]
        },
        {
            "num": "3", "title": "Strategy",
            "bullets": [
                ("Data Accuracy → ", "Integrate multi-satellite validation (Sentinel-2, Copernicus DEM) + local interpolation."),
                ("User Adoption → ", "Voice-assisted AI (SpringBot in tribal languages), minimal UI & offline-first sync."),
                ("Ground Truth Reliability → ", "Geo-tagged photo proof, auto-GPS timestamping & 2-tier administrative verification."),
                ("Regulatory Barriers → ", "Automated slope safety masking (>30° exclusion) & direct MGNREGA scheme alignment.")
            ]
        }
    ]

    sec_ys = [ly + Inches(0.2), ly + Inches(2.0), ly + Inches(3.8)]

    for idx, s_data in enumerate(sections_data):
        sy = sec_ys[idx]

        # Numbered Circle
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, lx + Inches(0.15), sy, Inches(0.6), Inches(0.6))
        circle.fill.solid()
        circle.fill.fore_color.rgb = WHITE
        circle.line.color.rgb = BLUE_TEXT
        circle.line.width = Pt(2)
        tf_c = circle.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.text = s_data["num"]
        p_c.font.name = 'Georgia'
        p_c.font.size = Pt(14)
        p_c.font.bold = True
        p_c.font.color.rgb = NAVY_TITLE
        p_c.alignment = PP_ALIGN.CENTER

        # Title Header Tag Box
        tag_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, lx + Inches(0.85), sy + Inches(0.08), Inches(1.8), Inches(0.44))
        tag_box.fill.solid()
        tag_box.fill.fore_color.rgb = DARK_BLUE_TAG
        tag_box.line.fill.background()
        tf_tag = tag_box.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = s_data["title"]
        p_tag.font.name = 'Georgia'
        p_tag.font.size = Pt(12)
        p_tag.font.bold = True
        p_tag.font.color.rgb = WHITE
        p_tag.alignment = PP_ALIGN.CENTER

        # Text Frame for Bullets
        content_box = slide.shapes.add_textbox(lx + Inches(0.85), sy + Inches(0.48), lw - Inches(1.0), Inches(1.3))
        tf_cont = content_box.text_frame
        tf_cont.word_wrap = True

        for b_idx, (b_lead, b_body) in enumerate(s_data["bullets"]):
            p_b = tf_cont.paragraphs[0] if b_idx == 0 else tf_cont.add_paragraph()
            p_b.text = "• "
            p_b.font.name = 'Arial'
            p_b.font.size = Pt(9)
            p_b.font.bold = True
            p_b.font.color.rgb = NAVY_TITLE

            r_lead = p_b.add_run()
            r_lead.text = b_lead
            r_lead.font.bold = True
            r_lead.font.color.rgb = TEXT_BLACK

            r_body = p_b.add_run()
            r_body.text = b_body
            r_body.font.bold = False
            r_body.font.color.rgb = GRAY_TEXT

    # -------------------------------------------------------------------------
    # 4. RIGHT SECTION: OUR FEASIBLE SOLUTION (8 SOLUTION PILLS)
    # -------------------------------------------------------------------------
    rx = Inches(8.8)
    rw = Inches(4.133)
    ry = Inches(0.95)

    # Title
    t_box = slide.shapes.add_textbox(rx, ry, rw, Inches(0.4))
    tf_tb = t_box.text_frame
    p_tb = tf_tb.paragraphs[0]
    p_tb.text = "Our Feasible Solution"
    p_tb.font.name = 'Georgia'
    p_tb.font.size = Pt(14)
    p_tb.font.bold = True
    p_tb.font.color.rgb = NAVY_TITLE
    p_tb.alignment = PP_ALIGN.CENTER

    solution_pills = [
        "Transparent MGNREGA itemized cost estimation.",
        "Map-based DEM input + auto satellite data fetch.",
        "Smart hydro-geological assessment with detailed PDF DPR.",
        "Verified government structure catalog (Check Dams, Trenches).",
        "End-to-end service flow — DEM Delineation → PRSI → DPR.",
        "Voice & Local Tribal Language support (SpringBot).",
        "Offline-first mobile GIS sync for remote tribal belts.",
        "Runoff & water storage capacity estimates (+45d baseflow)."
    ]

    py_start = ry + Inches(0.45)
    pill_h = Inches(0.52)
    pill_gap = Inches(0.1)

    for p_idx, pill_text in enumerate(solution_pills):
        py = py_start + p_idx * (pill_h + pill_gap)
        
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, py, rw, pill_h)
        pill.fill.solid()
        pill.fill.fore_color.rgb = LIGHT_GREEN_PILL
        pill.line.color.rgb = PILL_BORDER
        pill.line.width = Pt(1)

        tf_p = pill.text_frame
        tf_p.word_wrap = True
        p_p = tf_p.paragraphs[0]
        p_p.text = pill_text
        p_p.font.name = 'Arial'
        p_p.font.size = Pt(9.5)
        p_p.font.bold = True
        p_p.font.color.rgb = NAVY_TITLE
        p_p.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------------------
    # 5. BOTTOM URL LINK BANNER
    # -------------------------------------------------------------------------
    url_banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(6.6), Inches(12.533), Inches(0.35))
    url_banner.fill.solid()
    url_banner.fill.fore_color.rgb = RGBColor(224, 242, 254) # #E0F2FE
    url_banner.line.color.rgb = BLUE_TEXT
    url_banner.line.width = Pt(1)

    tf_url = url_banner.text_frame
    p_url = tf_url.paragraphs[0]
    p_url.text = "https://www.spring-ai.in/assessment"
    p_url.font.name = 'Arial'
    p_url.font.size = Pt(10)
    p_url.font.bold = True
    p_url.font.color.rgb = BLUE_TEXT
    p_url.font.underline = True
    p_url.alignment = PP_ALIGN.CENTER

    # 6. Footer Blue Bar
    footer_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.45))
    footer_bar.fill.solid()
    footer_bar.fill.fore_color.rgb = FOOTER_BLUE
    footer_bar.line.fill.background()

    tf_foot = footer_bar.text_frame
    p_foot_l = tf_foot.paragraphs[0]
    p_foot_l.text = "@SIH Idea submission- Template"
    p_foot_l.font.name = 'Arial'
    p_foot_l.font.size = Pt(11)
    p_foot_l.font.color.rgb = WHITE
    p_foot_l.alignment = PP_ALIGN.LEFT

    num_box = slide.shapes.add_textbox(Inches(12.2), Inches(7.05), Inches(1.0), Inches(0.45))
    tf_num = num_box.text_frame
    p_num = tf_num.paragraphs[0]
    p_num.text = "4"
    p_num.font.name = 'Arial'
    p_num.font.size = Pt(12)
    p_num.font.bold = True
    p_num.font.color.rgb = WHITE
    p_num.alignment = PP_ALIGN.RIGHT

    out_path = "SPRING_AI_Exact_Winning_Feasibility.pptx"
    prs.save(out_path)
    print(f"Successfully generated exact winning Feasibility slide at: {os.path.abspath(out_path)}")

if __name__ == "__main__":
    build_exact_winning_feasibility()
