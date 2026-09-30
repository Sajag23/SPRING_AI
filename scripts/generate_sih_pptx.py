"""
SPRING AI — Winning Grade SIH 2026 PowerPoint Generator
Generates a visually stunning 16:9 Widescreen Presentation Deck (.pptx)
Matching the official SIH 2026 Template Structure with high-impact visual design.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_sih_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # High-Impact Color System
    NAVY = RGBColor(15, 23, 42)         # #0F172A
    FOREST = RGBColor(8, 59, 44)        # #083B2C
    MINT = RGBColor(16, 185, 129)       # #10B981
    BLUE = RGBColor(37, 99, 235)        # #2563EB
    LIGHT_BG = RGBColor(248, 250, 252)  # #F8FAFC
    WHITE = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(203, 213, 225)# #CBD5E1
    RED_ACCENT = RGBColor(225, 29, 72)  # #E11D48
    SLATE_TEXT = RGBColor(71, 85, 105)  # #475569

    blank_layout = prs.slide_layouts[6]

    # Helper: Slide Header
    def add_header(slide, title_text, slide_num):
        # Top banner shape
        hdr_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        hdr_bg.fill.solid()
        hdr_bg.fill.fore_color.rgb = NAVY
        hdr_bg.line.fill.background()

        # Title
        tbox = slide.shapes.add_textbox(Inches(1.2), Inches(0.15), Inches(8.5), Inches(0.8))
        tf = tbox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = 'Space Grotesk'
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = WHITE

        # SIH 2026 Tag
        sih_box = slide.shapes.add_textbox(Inches(9.2), Inches(0.2), Inches(3.6), Inches(0.7))
        sih_tf = sih_box.text_frame
        p_sih = sih_tf.paragraphs[0]
        p_sih.text = f"SMART INDIA HACKATHON 2026  |  Slide 0{slide_num}"
        p_sih.font.name = 'Space Grotesk'
        p_sih.font.size = Pt(11)
        p_sih.font.bold = True
        p_sih.font.color.rgb = MINT
        p_sih.alignment = PP_ALIGN.RIGHT

        # Team Logo Circle
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.3), Inches(0.15), Inches(0.8), Inches(0.8))
        badge.fill.solid()
        badge.fill.fore_color.rgb = FOREST
        badge.line.color.rgb = MINT
        badge.line.width = Pt(1.5)
        btf = badge.text_frame
        bp = btf.paragraphs[0]
        bp.text = "SPRING\nAI"
        bp.font.size = Pt(8.5)
        bp.font.bold = True
        bp.font.color.rgb = WHITE
        bp.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 1: TITLE PAGE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Full background card
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Hero Banner Card inside
    hero1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.9))
    hero1.fill.solid()
    hero1.fill.fore_color.rgb = FOREST
    hero1.line.color.rgb = MINT
    hero1.line.width = Pt(2)
    htf = hero1.text_frame
    htf.word_wrap = True

    p1 = htf.paragraphs[0]
    p1.text = "SMART INDIA HACKATHON 2026"
    p1.font.name = 'Space Grotesk'
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = MINT
    p1.alignment = PP_ALIGN.CENTER

    p2 = htf.add_paragraph()
    p2.text = "MINISTRY OF TRIBAL AFFAIRS • OFFICIAL SUBMISSION"
    p2.font.name = 'Space Grotesk'
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER

    # Blank line
    htf.add_paragraph()

    # Key metadata table items
    metadata = [
        ("Problem Statement ID:", "SIH1689 (Ministry of Tribal Affairs)"),
        ("Problem Statement Title:", "AI-Powered Spring Revival & Recharge Planning System"),
        ("Theme & Category:", "Agriculture, Food & Rural Technology | Software"),
        ("Team Name:", "SPRING AI"),
        ("Core Capability:", "Springshed Delineation • PRSI AI Scoring • Slope >30° Mask • MGNREGA PDF Dossiers")
    ]

    for label, val in metadata:
        p = htf.add_paragraph()
        p.text = f"  📍 {label}  {val}"
        p.font.name = 'Plus Jakarta Sans'
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 2: IDEA TITLE
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "IDEA TITLE", 2)

    # Banner
    b2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.65))
    b2.fill.solid()
    b2.fill.fore_color.rgb = FOREST
    b2.line.color.rgb = MINT
    b2_tf = b2.text_frame
    pb2 = b2_tf.paragraphs[0]
    pb2.text = "SPRING AI — Intelligent Spring Revival & Recharge Planning Platform"
    pb2.font.name = 'Space Grotesk'
    pb2.font.size = Pt(18)
    pb2.font.bold = True
    pb2.font.color.rgb = WHITE
    pb2.alignment = PP_ALIGN.CENTER

    # 3 Columns
    col_w = Inches(3.7)
    gap = Inches(0.3)
    start_x = Inches(0.8)

    boxes2 = [
        ("THE GAP (CHALLENGE)", RED_ACCENT, [
            "Data fragmented across elevation, rainfall & geology.",
            "Subsurface springshed boundaries hard to trace manually.",
            "Drying natural springs cause severe summer water crisis.",
            "Unguided excavation on steep slopes induces landslides."
        ]),
        ("THE SOLUTION (OUR AI)", FOREST, [
            "Fuse DEM 30m, IMD rainfall & GSI lithology.",
            "Delineate springsheds & calculate suitability score (PRSI).",
            "Recommend structures (Check Dams, Trenches, Ponds).",
            "Automated slope >30° terrain safety exclusion mask.",
            "Generate MGNREGA budget & persondays."
        ]),
        ("THE VALUE (OUTCOMES)", BLUE, [
            "Restores springs in 100+ Scheduled Tribal Districts.",
            "Increases groundwater table (+1.2m to +2.5m).",
            "Extends summer baseflow by +45 days.",
            "Optimizes MGNREGA budget & local labor.",
            "Produces 1-click official PDF dossiers."
        ])
    ]

    for i, (title, color, items) in enumerate(boxes2):
        x = start_x + i * (col_w + gap)
        box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.1), col_w, Inches(4.1))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = color
        box.line.width = Pt(2)
        btf = box.text_frame
        btf.word_wrap = True

        pt = btf.paragraphs[0]
        pt.text = title
        pt.font.name = 'Space Grotesk'
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = color

        for item in items:
            pi = btf.add_paragraph()
            pi.text = f"• {item}"
            pi.font.name = 'Plus Jakarta Sans'
            pi.font.size = Pt(11)
            pi.font.color.rgb = NAVY

    # Innovation Banner
    inno = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.75))
    inno.fill.solid()
    inno.fill.fore_color.rgb = LIGHT_BG
    inno.line.color.rgb = BLUE
    inno_tf = inno.text_frame
    inno_tf.word_wrap = True
    pi = inno_tf.paragraphs[0]
    pi.text = "INNOVATION: Multi-spatial satellite fusion + Ensemble ML (RF + XGBoost, F1: 0.9555) + Slope >30° Landslide Safety Mask + MGNREGA Costing Engine + 1-Click PDF Dossier Exporter"
    pi.font.name = 'Plus Jakarta Sans'
    pi.font.size = Pt(11.5)
    pi.font.bold = True
    pi.font.color.rgb = NAVY
    pi.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "TECHNICAL APPROACH", 3)

    # 6 Pipeline Blocks
    pipe_blocks = [
        ("01 INGEST", "Copernicus DEM\nIMD Rainfall\nGSI Lithology"),
        ("02 FUSE", "Slope Extractor\nRainfall Resampler\nLineament Parser"),
        ("03 STORE", "SQLite / GeoJSON\nSpring Catchments\nGround Logs"),
        ("04 GEOAI CORE", "Random Forest\nXGBoost Ensemble\nPRSI Scoring"),
        ("05 SAFETY", "Slope >30° Mask\nLandslide Exclusion\nVegetative Plan"),
        ("06 PORTAL", "Folium GIS Canvas\nMGNREGA Costing\nPDF Dossier Exporter")
    ]

    bw = Inches(1.75)
    bgap = Inches(0.2)
    by = Inches(1.3)

    for i, (btitle, bdesc) in enumerate(pipe_blocks):
        bx = Inches(0.8) + i * (bw + bgap)
        bbox = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, bw, Inches(2.2))
        bbox.fill.solid()
        bbox.fill.fore_color.rgb = WHITE
        bbox.line.color.rgb = BLUE
        bbox.line.width = Pt(1.5)
        btf = bbox.text_frame
        btf.word_wrap = True

        bp1 = btf.paragraphs[0]
        bp1.text = btitle
        bp1.font.name = 'Space Grotesk'
        bp1.font.size = Pt(11)
        bp1.font.bold = True
        bp1.font.color.rgb = BLUE
        bp1.alignment = PP_ALIGN.CENTER

        bp2 = btf.add_paragraph()
        bp2.text = bdesc
        bp2.font.name = 'Plus Jakarta Sans'
        bp2.font.size = Pt(10)
        bp2.font.color.rgb = NAVY
        bp2.alignment = PP_ALIGN.CENTER

    # Lower Boxes
    ts_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.7), Inches(5.7), Inches(3.4))
    ts_box.fill.solid()
    ts_box.fill.fore_color.rgb = LIGHT_BG
    ts_box.line.color.rgb = BLUE
    ts_tf = ts_box.text_frame
    ts_tf.word_wrap = True

    pts = ts_tf.paragraphs[0]
    pts.text = "TECH STACK & INFRASTRUCTURE"
    pts.font.name = 'Space Grotesk'
    pts.font.size = Pt(14)
    pts.font.bold = True
    pts.font.color.rgb = BLUE

    ts_list = [
        "Language & UI: Python 3.10+ • Streamlit Framework",
        "GIS Canvas: Folium • Leaflet.js • Esri Satellite HD",
        "Machine Learning: Scikit-Learn • XGBoost (F1: 0.9555)",
        "Spatial Data: GeoPandas • Copernicus 30m DEM",
        "PDF Exporter: FPDF2 Engine for 1-Click Dossiers"
    ]
    for item in ts_list:
        p = ts_tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = 'Plus Jakarta Sans'
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY

    al_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.7), Inches(5.7), Inches(3.4))
    al_box.fill.solid()
    al_box.fill.fore_color.rgb = LIGHT_BG
    al_box.line.color.rgb = FOREST
    al_tf = al_box.text_frame
    al_tf.word_wrap = True

    pal = al_tf.paragraphs[0]
    pal.text = "ANALYTICS LOGIC & MATH FORMULA"
    pal.font.name = 'Space Grotesk'
    pal.font.size = Pt(14)
    pal.font.bold = True
    pal.font.color.rgb = FOREST

    al_list = [
        "PRSI Formula: 0.32(Rainfall) + 0.24(Geology) + 0.20(Slope) + 0.14(Fractures) + 0.06(LULC) + 0.04(Soil)",
        "Safety Rule: If Slope >30° => PRSI = 0 (Landslide Exclusion)",
        "Water Impact: Volume = Treated Catchment Area × Infiltration",
        "MGNREGA Costing: Persondays = Total Cost / Labor Rate"
    ]
    for item in al_list:
        p = al_tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = 'Plus Jakarta Sans'
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY

    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "FEASIBILITY AND VIABILITY", 4)

    feas_boxes = [
        ("FEASIBILITY (PROOF)", FOREST, [
            "✓ Full-stack prototype live at http://localhost:8501",
            "✓ Pre-configured with 10 Real Hydro Baselines",
            "✓ Automated Terrain Safety Masking (Slope >30°)",
            "✓ MGNREGA Budget & Persondays Costing Engine",
            "✓ Automated 18-step Unit Test Suite (100% Pass)"
        ]),
        ("KEY RISKS", RED_ACCENT, [
            "Δ Data quality & missing ground logs in remote hills",
            "Δ Uncertainty in subsurface rock fracture lineaments",
            "Δ Landslide hazards caused by unguided digging",
            "Δ Low digital literacy among field officers",
            "Δ Scale of real-world multi-state data aggregation"
        ]),
        ("MITIGATION STRATEGY", BLUE, [
            "→ Fuse Copernicus DEM 30m, NASA & IMD APIs",
            "→ Confidence scoring & uncertainty indicators",
            "→ Automated Slope >30° Landslide Masking",
            "→ Minimalist 1-click UI & 1-click PDF Dossier",
            "→ Modular python architecture & memory caching"
        ])
    ]

    for i, (title, color, items) in enumerate(feas_boxes):
        x = start_x + i * (col_w + gap)
        box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.4), col_w, Inches(4.7))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = color
        box.line.width = Pt(2)
        btf = box.text_frame
        btf.word_wrap = True

        pt = btf.paragraphs[0]
        pt.text = title
        pt.font.name = 'Space Grotesk'
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = color

        for item in items:
            pi = btf.add_paragraph()
            pi.text = item
            pi.font.name = 'Plus Jakarta Sans'
            pi.font.size = Pt(11)
            pi.font.color.rgb = NAVY

    # Path Banner
    path_b = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.75))
    path_b.fill.solid()
    path_b.fill.fore_color.rgb = NAVY
    path_b.line.fill.background()
    p_tf = path_b.text_frame
    ppt = p_tf.paragraphs[0]
    ppt.text = "VIABLE MVP PATH: Ingest Satellite Data → Delineate Springshed → Score PRSI → Mask Slope >30° → Generate PDF Dossier"
    ppt.font.name = 'Plus Jakarta Sans'
    ppt.font.size = Pt(12.5)
    ppt.font.bold = True
    ppt.font.color.rgb = WHITE
    ppt.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "IMPACT AND BENEFITS", 5)

    # Central Box
    cb = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(2.9), Inches(4.3), Inches(1.8))
    cb.fill.solid()
    cb.fill.fore_color.rgb = NAVY
    cb.line.color.rgb = MINT
    cb.line.width = Pt(2)
    cbtf = cb.text_frame
    pcb = cbtf.paragraphs[0]
    pcb.text = "SPRING AI\nEVIDENCE → RECHARGE\nINTELLIGENCE"
    pcb.font.name = 'Space Grotesk'
    pcb.font.size = Pt(16)
    pcb.font.bold = True
    pcb.font.color.rgb = WHITE
    pcb.alignment = PP_ALIGN.CENTER

    # 4 Corner Boxes
    corners = [
        ("MINISTRY & POLICY MAKERS", ["• Data-driven decisions across 100+ districts", "• MGNREGA budget optimization", "• National water security monitoring"], Inches(0.8), Inches(1.4)),
        ("FIELD SURVEYORS & ENGINEERS", ["• Site structure recommendations", "• Automated MGNREGA persondays", "• 1-Click official PDF Dossiers"], Inches(8.3), Inches(1.4)),
        ("HYDROLOGISTS & SCIENTISTS", ["• Subsurface springshed boundary delineation", "• Transparent GeoAI feature weights (RF+XGBoost)", "• Ground survey re-scoring loop"], Inches(0.8), Inches(4.1)),
        ("TRIBAL COMMUNITIES & VILLAGES", ["• 18.5 Billion L annual water storage unlocked", "• Extended summer baseflow (+45 days)", "• Livelihood security for tribal households"], Inches(8.3), Inches(4.1))
    ]

    for title, items, x, y in corners:
        box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(4.2), Inches(2.3))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BLUE
        box.line.width = Pt(1.5)
        btf = box.text_frame
        btf.word_wrap = True

        pt = btf.paragraphs[0]
        pt.text = title
        pt.font.name = 'Space Grotesk'
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = BLUE

        for item in items:
            pi = btf.add_paragraph()
            pi.text = item
            pi.font.name = 'Plus Jakarta Sans'
            pi.font.size = Pt(10.5)
            pi.font.color.rgb = NAVY

    # Demo Banner
    demo_b = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.6))
    demo_b.fill.solid()
    demo_b.fill.fore_color.rgb = LIGHT_BG
    demo_b.line.color.rgb = BLUE
    dtf = demo_b.text_frame
    pdt = dtf.paragraphs[0]
    pdt.text = "DEMONSTRATION VALUE: Converts fragmented satellite elevation, geology, rainfall, and ground logs into a single, visual decision picture."
    pdt.font.name = 'Plus Jakarta Sans'
    pdt.font.size = Pt(11)
    pdt.font.bold = True
    pdt.font.color.rgb = NAVY
    pdt.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "RESEARCH AND REFERENCES", 6)

    rw = Inches(5.7)
    rh = Inches(4.7)

    ref1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), rw, rh)
    ref1.fill.solid()
    ref1.fill.fore_color.rgb = WHITE
    ref1.line.color.rgb = BLUE
    ref1.line.width = Pt(1.5)
    r1tf = ref1.text_frame
    r1tf.word_wrap = True

    pr1 = r1tf.paragraphs[0]
    pr1.text = "CORE REFERENCES & TECHNOLOGY"
    pr1.font.name = 'Space Grotesk'
    pr1.font.size = Pt(15)
    pr1.font.bold = True
    pr1.font.color.rgb = BLUE

    tech_refs = [
        "Streamlit & Python 3.10+ — Web UI framework & data pipeline.",
        "Folium / Leaflet.js — Interactive GIS mapping & satellite rendering.",
        "Scikit-Learn & XGBoost — Supervised ensemble ML classifiers.",
        "FPDF2 Engine — Automated PDF document & executive dossier generation.",
        "Copernicus & Bhuvan 30m DEM — Elevation, slope (°), and aspect topography.",
        "IMD Pune & GSI Bhukosh — Rainfall grids & lithological rock formations.",
        "CGWB NAQUIM — National Aquifer Mapping hydrogeological baselines."
    ]
    for item in tech_refs:
        p = r1tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = 'Plus Jakarta Sans'
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY

    ref2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.4), rw, rh)
    ref2.fill.solid()
    ref2.fill.fore_color.rgb = WHITE
    ref2.line.color.rgb = FOREST
    ref2.line.width = Pt(1.5)
    r2tf = ref2.text_frame
    r2tf.word_wrap = True

    pr2 = r2tf.paragraphs[0]
    pr2.text = "PROJECT-DERIVED RESEARCH BASIS"
    pr2.font.name = 'Space Grotesk'
    pr2.font.size = Pt(15)
    pr2.font.bold = True
    pr2.font.color.rgb = FOREST

    res_refs = [
        "PRSI Index Formulation → Multi-criteria geospatial suitability overlay.",
        "Random Forest + XGBoost → Supervised training on CGWB baselines (F1: 0.9555).",
        "DEM Hydrologic Flow Direction → Delineates catchment area & stream order.",
        "Slope Hazard Rule → Slopes >30° strictly masked from heavy excavation.",
        "MGNREGA Costing Formula → Links structure volume to persondays labor rates."
    ]
    for item in res_refs:
        p = r2tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = 'Plus Jakarta Sans'
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY

    src_box = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.75))
    src_box.fill.solid()
    src_box.fill.fore_color.rgb = NAVY
    src_box.line.fill.background()
    stf = src_box.text_frame
    pst = stf.paragraphs[0]
    pst.text = "SOURCE: SPRING AI README • Architecture, algorithms, APIs, testing & implementation details documented in submitted project repository."
    pst.font.name = 'Plus Jakarta Sans'
    pst.font.size = Pt(11.5)
    pst.font.bold = True
    pst.font.color.rgb = WHITE
    pst.alignment = PP_ALIGN.CENTER

    out_path = "SPRING_AI_SIH_2026_Winner_Deck.pptx"
    prs.save(out_path)
    print(f"Successfully updated PowerPoint presentation at {out_path}")

if __name__ == "__main__":
    build_sih_presentation()
